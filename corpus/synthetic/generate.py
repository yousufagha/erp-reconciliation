"""Generates the synthetic scenario pack: one fabricator's August 2026 paperwork.

Every business, person, ABN and number here is invented. ABNs are generated to
pass the ABN checksum so validation can be exercised; they belong to no one.

The generator is standalone on purpose: it never imports the `pile` package, so
the test documents are not shaped by the code under test. Ground truth is
computed from the scenario below, not read back from the rendered files.

Usage: python corpus/synthetic/generate.py            (writes into ./scenario_aug2026)
"""
from __future__ import annotations

import csv
import io
import json
import random
import shutil
from dataclasses import dataclass, field
from datetime import date, timedelta
from email.message import EmailMessage
from pathlib import Path

from PIL import Image, ImageFilter, ImageOps
import pypdfium2 as pdfium
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

OUT = Path(__file__).resolve().parent / "scenario_aug2026"
RNG = random.Random(20260826)
GST_RATE = 0.10


# ============================================================== scenario data
def make_abn(rng: random.Random) -> str:
    weights = [10, 1, 3, 5, 7, 9, 11, 13, 15, 17, 19]
    while True:
        d = [rng.randint(1, 9)] + [rng.randint(0, 9) for _ in range(10)]
        s = sum(w * (x - (1 if i == 0 else 0)) for i, (w, x) in enumerate(zip(weights, d)))
        if s % 89 == 0:
            n = "".join(map(str, d))
            return f"{n[:2]} {n[2:5]} {n[5:8]} {n[8:]}"


@dataclass
class Party:
    key: str
    name: str
    abn: str
    address: list[str]
    email: str


@dataclass
class Item:
    sku: str
    description: str
    unit: str


BUYER = Party("BUY", "Ridgeline Steel Fabrications Pty Ltd", make_abn(RNG),
              ["Unit 4, 18 Anvil Court", "Dandenong South VIC 3175"], "accounts@ridgeline-fab.example")
SUP = {
    "S1": Party("S1", "Harbourline Steel Supply Pty Ltd", make_abn(RNG),
                ["220 Wharf Road", "Port Melbourne VIC 3207"], "ar@harbourline-steel.example"),
    "S2": Party("S2", "Boltmaster Fasteners Pty Ltd", make_abn(RNG),
                ["9 Rivet Street", "Campbellfield VIC 3061"], "invoices@boltmaster.example"),
    "S3": Party("S3", "Coastal Coatings Pty Ltd", make_abn(RNG),
                ["51 Enamel Drive", "Geelong VIC 3220"], "accounts@coastalcoatings.example"),
    "S4": Party("S4", "Southern Gas & Welding Supplies Pty Ltd", make_abn(RNG),
                ["3/77 Arc Lane", "Moorabbin VIC 3189"], "billing@southerngasweld.example"),
}
ITEMS = {
    "UB310": Item("STL-UB-310-46", "UB 310x46 Grade 300PLUS 9.0m", "t"),
    "RHS100": Item("STL-RHS-100-5G", "RHS 100x100x5.0 Galvanised", "t"),
    "PL12": Item("STL-PL-12-300", "Plate 12mm Grade 300PLUS", "t"),
    "FB75": Item("STL-FB-75-10", "Flat Bar 75x10 Grade 300", "m"),
    "M16": Item("BF-M16-HDG-SET", "Bolt set M16 x 50 HDG with nut and washer", "ea"),
    "M20A": Item("BF-M20-CHEM", "Chemical anchor M20 x 260", "ea"),
    "N16": Item("BF-N16-BOX", "Hex nut M16 HDG box of 100", "box"),
    "ZP20": Item("CC-ZP-20L", "Inorganic zinc primer 20L pail", "pail"),
    "EP16": Item("CC-EP-16K", "Epoxy topcoat 16L kit grey", "kit"),
    "GAS": Item("SG-MIX-G", "Mixed shielding gas G size cylinder refill", "ea"),
    "WIRE": Item("SG-ER70-09", "Welding wire ER70S-6 0.9mm 15kg spool", "spool"),
}
SUPPLIER_OF = {"UB310": "S1", "RHS100": "S1", "PL12": "S1", "FB75": "S1", "M16": "S2", "M20A": "S2",
               "N16": "S2", "ZP20": "S3", "EP16": "S3", "GAS": "S4", "WIRE": "S4"}
SCHEDULE_S1 = {"UB310": 1840.00, "RHS100": 2310.00, "PL12": 1690.00, "FB75": 18.40}

D = lambda day: date(2026, 8, day)  # noqa: E731


@dataclass
class Line:
    item: str
    qty: float
    price: float | None = None

    @property
    def amount(self) -> float:
        return round(self.qty * (self.price or 0), 2)


@dataclass
class PO:
    number: str
    supplier: str
    date: date
    deliver_by: date
    lines: list[Line]


@dataclass
class Receipt:
    number: str
    kind: str            # "docket" (supplier's) or "grn" (buyer's)
    supplier: str
    po: str
    date: date
    lines: list[Line]    # qty received


@dataclass
class Invoice:
    number: str
    supplier: str
    date: date
    po: str | None       # PO number printed on the invoice (None = not quoted)
    lines: list[Line]
    terms: str = "30 days from invoice"
    matches_po: str | None = None   # which PO it really belongs to (truth)
    credit_for: str | None = None   # credit notes: original invoice

    @property
    def subtotal(self) -> float:
        return round(sum(l.amount for l in self.lines), 2)

    @property
    def gst(self) -> float:
        return round(self.subtotal * GST_RATE, 2)

    @property
    def total(self) -> float:
        return round(self.subtotal + self.gst, 2)


@dataclass
class Payment:
    number: str
    supplier: str
    date: date
    allocations: list[tuple[str, float]]  # (invoice number, amount paid)

    @property
    def total(self) -> float:
        return round(sum(a for _, a in self.allocations), 2)


POS = [
    PO("PO-4101", "S1", D(3), D(7), [Line("UB310", 4.990, 1840.00), Line("RHS100", 0.860, 2310.00), Line("PL12", 1.240, 1690.00)]),
    PO("PO-4102", "S1", D(5), D(11), [Line("FB75", 120, 18.40), Line("RHS100", 1.200, 2310.00)]),
    PO("PO-4103", "S2", D(6), D(8), [Line("M16", 400, 4.00), Line("M20A", 150, 7.85), Line("N16", 20, 32.50)]),
    PO("PO-4104", "S2", D(10), D(21), [Line("M16", 600, 4.00), Line("M20A", 200, 7.85)]),
    PO("PO-4105", "S3", D(11), D(25), [Line("ZP20", 12, 289.00), Line("EP16", 6, 412.50)]),
    PO("PO-4106", "S3", D(12), D(19), [Line("ZP20", 8, 289.00)]),
    PO("PO-4107", "S4", D(13), D(15), [Line("GAS", 10, 96.00), Line("WIRE", 6, 138.00)]),
    PO("PO-4108", "S4", D(20), D(31), [Line("WIRE", 4, 138.00)]),
]
RECEIPTS = [
    Receipt("DD-55120", "docket", "S1", "PO-4101", D(7), [Line("UB310", 4.990), Line("RHS100", 0.860), Line("PL12", 1.240)]),
    Receipt("GRN-7002", "grn", "S1", "PO-4102", D(11), [Line("FB75", 120), Line("RHS100", 0.960)]),
    Receipt("BF-20931", "docket", "S2", "PO-4103", D(8), [Line("M16", 400), Line("M20A", 150), Line("N16", 20)]),
    Receipt("BF-21007", "docket", "S2", "PO-4104", D(14), [Line("M16", 600)]),
    Receipt("BF-21088", "docket", "S2", "PO-4104", D(21), [Line("M20A", 200)]),
    Receipt("GRN-7006", "grn", "S3", "PO-4106", D(19), [Line("ZP20", 8)]),
    Receipt("SG-3301", "docket", "S4", "PO-4107", D(15), [Line("GAS", 10), Line("WIRE", 6)]),
]
INVOICES = [
    Invoice("HS-88214", "S1", D(7), "PO-4101", [Line("UB310", 4.990, 1840.00), Line("RHS100", 0.860, 2310.00), Line("PL12", 1.240, 1690.00)], matches_po="PO-4101"),
    Invoice("HS-88302", "S1", D(12), "PO-4102", [Line("FB75", 120, 18.40), Line("RHS100", 1.200, 2310.00)], matches_po="PO-4102"),
    Invoice("BF-INV-30551", "S2", D(9), "PO-4103", [Line("M16", 400, 4.20), Line("M20A", 150, 7.85), Line("N16", 20, 32.50)], terms="Net 30", matches_po="PO-4103"),
    Invoice("BF-INV-30602", "S2", D(14), "PO-4104", [Line("M16", 600, 4.00)], terms="Net 30", matches_po="PO-4104"),
    Invoice("BF-INV-30688", "S2", D(24), "PO-4104", [Line("M20A", 200, 7.85)], terms="Net 30", matches_po="PO-4104"),
    Invoice("CC-10442", "S3", D(15), "PO-4105", [Line("ZP20", 12, 289.00), Line("EP16", 6, 412.50)], terms="14 days", matches_po="PO-4105"),
    Invoice("SG-77120", "S4", D(18), None, [Line("GAS", 10, 96.00), Line("WIRE", 6, 138.00)], terms="30 days EOM", matches_po="PO-4107"),
    Invoice("CC-10497", "S3", D(22), None, [], terms="14 days", matches_po=None),
]
# a service line with no item-master entry (no PO exists for it)
INVOICES[-1].lines = [Line("SERVICE", 1, 640.00)]
ITEMS["SERVICE"] = Item("CC-SVC-FILT", "Spray booth filter replacement (service call)", "ea")
SUPPLIER_OF["SERVICE"] = "S3"

CREDITS = [
    Invoice("HS-CN-1190", "S1", D(20), "PO-4102", [Line("RHS100", 0.240, 2310.00)], credit_for="HS-88302", matches_po="PO-4102"),
]
# An invoice the supplier issued that never reached the pile. It exists only on their statement.
UNSEEN = Invoice("HS-88377", "S1", D(25), "PO-4109", [Line("PL12", 0.500, 1690.00)])

PAYMENTS = [
    Payment("RA-9001", "S1", D(29), [("HS-88214", INVOICES[0].total)]),
    Payment("RA-9002", "S2", D(29), [("BF-INV-30602", INVOICES[3].total)]),
    Payment("RA-9003", "S4", D(29), [("SG-77120", 1000.00)]),
]


# ============================================================== formatting helpers
def money(x: float) -> str:
    return f"{x:,.2f}"


def qty(x: float, unit: str) -> str:
    return f"{x:,.3f}" if unit == "t" else (f"{x:,.0f}" if float(x).is_integer() else f"{x:,.2f}")


DATE_STYLES = {
    "S1": lambda d: d.strftime("%d/%m/%Y"),
    "S2": lambda d: d.strftime("%-d %b %Y"),
    "S3": lambda d: d.strftime("%d.%m.%Y"),
    "S4": lambda d: d.isoformat(),
    "BUY": lambda d: d.strftime("%d/%m/%Y"),
}

STYLES = getSampleStyleSheet()
SMALL = ParagraphStyle("small", parent=STYLES["Normal"], fontSize=8.5, leading=10.5)
NORMAL = ParagraphStyle("norm", parent=STYLES["Normal"], fontSize=9.5, leading=12)
TITLE = ParagraphStyle("title", parent=STYLES["Title"], fontSize=18, leading=22, alignment=2)
TITLE_C = ParagraphStyle("titlec", parent=TITLE, alignment=1)
BIG = ParagraphStyle("big", parent=STYLES["Heading2"], fontSize=13, leading=16)


def grid(rows, widths, header_bg=colors.HexColor("#e6e9ec"), align_right_from=None, box=True):
    t = Table(rows, colWidths=widths, repeatRows=1)
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", 8.8),
          ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.8),
          ("BACKGROUND", (0, 0), (-1, 0), header_bg),
          ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3)]
    if box:
        st += [("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.grey)]
    if align_right_from is not None:
        st.append(("ALIGN", (align_right_from, 0), (-1, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def kv(rows, widths=(32 * mm, 45 * mm)):
    t = Table(rows, colWidths=widths)
    t.setStyle(TableStyle([("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9),
                           ("FONT", (1, 0), (1, -1), "Helvetica", 9),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("TOPPADDING", (0, 0), (-1, -1), 1.5)]))
    return t


def party_block(p: Party, with_abn=True):
    lines = [f"<b>{p.name}</b>"] + p.address + ([f"ABN {p.abn}"] if with_abn else []) + [p.email]
    return Paragraph("<br/>".join(lines), NORMAL)


def two_col(left, right, widths=(95 * mm, 80 * mm)):
    t = Table([[left, right]], colWidths=widths)
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def build_pdf(path: Path, pages: list[list]) -> None:
    """pages: list of flowable lists; a page break goes between them."""
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title=path.stem,
                            author="synthetic", creator="synthetic generator")
    story = []
    for i, p in enumerate(pages):
        if i:
            story.append(PageBreak())
        story += p
    doc.build(story)


# ============================================================== document renderers
def invoice_story(inv: Invoice, credit=False, page_split=None, copy_stamp=False):
    """Returns a list of pages. Layout depends on the supplier (three house styles)."""
    s = SUP[inv.supplier]
    fmt = DATE_STYLES[inv.supplier]
    due = inv.date + timedelta(days=30 if "30" in inv.terms else 14)
    lines = inv.lines
    if s.key in ("S1",) or inv.number == "CC-10497":
        title = "TAX CREDIT NOTE" if credit else "TAX INVOICE"
        head = kv([["Credit Note No" if credit else "Invoice No", inv.number],
                   ["Date", fmt(inv.date)],
                   ["Customer PO", inv.po or ""],
                   ["Original Invoice" if credit else "Due Date", inv.credit_for if credit else fmt(due)],
                   ["Account", "RIDG01"]])
        top = [two_col(party_block(s), [Paragraph(title, TITLE), Spacer(1, 4), head])]
        bill = [Spacer(1, 8), Paragraph("<b>Bill To</b><br/>" + BUYER.name + "<br/>" + "<br/>".join(BUYER.address) +
                                        f"<br/>ABN {BUYER.abn}", NORMAL), Spacer(1, 10)]
        rows = [["Code", "Description", "Qty", "Unit", "Unit Price", "Amount"]]
        for l in lines:
            it = ITEMS[l.item]
            rows.append([it.sku, Paragraph(it.description, SMALL), qty(l.qty, it.unit), it.unit, money(l.price), money(l.amount)])
        widths = [30 * mm, 62 * mm, 17 * mm, 11 * mm, 22 * mm, 24 * mm]
        tot = kv([["Subtotal", money(inv.subtotal)], ["GST 10%", money(inv.gst)],
                  ["Total Credit AUD" if credit else "Total AUD", money(inv.total)]], widths=(35 * mm, 30 * mm))
        foot = [Spacer(1, 10), Paragraph(f"Terms: {inv.terms}. EFT to BSB 063-000 Acct 1234 5678, ref {inv.number}. "
                                         "All goods remain the property of the seller until paid in full.", SMALL)]
    elif s.key == "S2":
        title = "Tax Invoice"
        head = kv([["Invoice #", inv.number], ["Invoice Date", fmt(inv.date)], ["Your Order", inv.po or ""],
                   ["Terms", inv.terms]], widths=(28 * mm, 40 * mm))
        top = [Paragraph(title, TITLE_C), Spacer(1, 6), two_col(party_block(s), head, widths=(100 * mm, 75 * mm))]
        bill = [Spacer(1, 8), Paragraph("Deliver to / Invoice to: <b>" + BUYER.name + "</b>, " + ", ".join(BUYER.address), NORMAL), Spacer(1, 10)]
        rows = [["Qty", "Item", "Description", "Price Ex GST", "Line Total"]]
        for l in lines:
            it = ITEMS[l.item]
            rows.append([qty(l.qty, it.unit), it.sku, Paragraph(it.description, SMALL), money(l.price), money(l.amount)])
        widths = [16 * mm, 34 * mm, 72 * mm, 22 * mm, 24 * mm]
        tot = kv([["Total ex GST", money(inv.subtotal)], ["GST", money(inv.gst)], ["Amount Due", "$" + money(inv.total)]],
                 widths=(35 * mm, 30 * mm))
        foot = [Spacer(1, 10), Paragraph("Please quote the invoice number with your payment. Thank you for your business.", SMALL)]
    else:  # S3 / S4 compact style
        title = "Tax Invoice"
        head = kv([["Our Ref", inv.number], ["Date", fmt(inv.date)], ["Order No", inv.po or "-"], ["Terms", inv.terms]],
                  widths=(22 * mm, 40 * mm))
        top = [two_col(Paragraph(f"<font size=15><b>{s.name}</b></font><br/>" + " | ".join(s.address) +
                                 f"<br/>A.B.N. {s.abn}", NORMAL), [Paragraph(title, BIG), head], widths=(110 * mm, 65 * mm))]
        bill = [Spacer(1, 8), Paragraph("Customer: " + BUYER.name, NORMAL), Spacer(1, 8)]
        rows = [["Description", "Quantity", "Rate", "Value"]]
        for l in lines:
            it = ITEMS[l.item]
            rows.append([Paragraph(f"{it.description} [{it.sku}]", SMALL), qty(l.qty, it.unit), money(l.price), money(l.amount)])
        widths = [100 * mm, 22 * mm, 24 * mm, 26 * mm]
        tot = kv([["Sub Total", money(inv.subtotal)], ["GST", money(inv.gst)], ["TOTAL INC GST", money(inv.total)]],
                 widths=(35 * mm, 30 * mm))
        foot = [Spacer(1, 10), Paragraph("Pay by EFT within terms. Queries to " + s.email, SMALL)]

    stamp = [Paragraph("<font color='#b00000' size=14><b>COPY</b></font>", NORMAL)] if copy_stamp else []
    table = grid(rows, widths, align_right_from=len(rows[0]) - 2)
    totals_right = Table([["", tot]], colWidths=[95 * mm, 80 * mm])
    if page_split:
        n = page_split
        first_rows = rows[: n + 1]
        rest_rows = [rows[0]] + rows[n + 1:]
        p1 = stamp + top + bill + [grid(first_rows, widths, align_right_from=len(rows[0]) - 2), Spacer(1, 12),
                                   Paragraph("Continued overleaf. Page 1 of 2", SMALL)]
        p2 = [Paragraph(f"{s.name} - {title} {inv.number} (continued)", NORMAL), Spacer(1, 8),
              grid(rest_rows, widths, align_right_from=len(rows[0]) - 2), Spacer(1, 10), totals_right] + foot + \
             [Spacer(1, 12), Paragraph("Page 2 of 2", SMALL)]
        return [p1, p2]
    return [stamp + top + bill + [table, Spacer(1, 10), totals_right] + foot]


def receipt_story(r: Receipt):
    s = SUP[r.supplier]
    po = next(p for p in POS if p.number == r.po)
    if r.kind == "docket":
        fmt = DATE_STYLES[r.supplier]
        head = kv([["Docket No", r.number], ["Date", fmt(r.date)], ["Your Order No", r.po], ["Vehicle", "Truck 3"]])
        top = [two_col(party_block(s, with_abn=False), [Paragraph("DELIVERY DOCKET", TITLE), Spacer(1, 4), head])]
        rows = [["Code", "Description", "Qty Supplied", "Unit"]]
        for l in r.lines:
            it = ITEMS[l.item]
            rows.append([it.sku, Paragraph(it.description, SMALL), qty(l.qty, it.unit), it.unit])
        body = [Spacer(1, 8), Paragraph("Deliver to: " + BUYER.name + ", " + ", ".join(BUYER.address), NORMAL),
                Spacer(1, 10), grid(rows, [34 * mm, 90 * mm, 26 * mm, 20 * mm], align_right_from=2),
                Spacer(1, 30), Paragraph("Received in good order by: ______________________  Signature: ____________", NORMAL),
                Spacer(1, 6), Paragraph("No prices shown. This is not a tax invoice.", SMALL)]
        return [top + body]
    head = kv([["GRN No", r.number], ["Received", DATE_STYLES["BUY"](r.date)], ["PO Number", r.po], ["Supplier", s.name]],
              widths=(28 * mm, 70 * mm))
    top = [two_col(party_block(BUYER), [Paragraph("GOODS RECEIVED NOTE", BIG), head], widths=(80 * mm, 100 * mm))]
    ordered = {l.item: l.qty for l in po.lines}
    rows = [["Item", "Description", "Ordered", "Received", "Unit"]]
    for l in r.lines:
        it = ITEMS[l.item]
        rows.append([it.sku, Paragraph(it.description, SMALL), qty(ordered[l.item], it.unit), qty(l.qty, it.unit), it.unit])
    body = [Spacer(1, 12), grid(rows, [34 * mm, 80 * mm, 22 * mm, 22 * mm, 16 * mm], align_right_from=2),
            Spacer(1, 12), Paragraph("Checked by: M. Okafor (stores)", NORMAL)]
    return [top + body]


def po_story(po: PO):
    s = SUP[po.supplier]
    head = kv([["PO Number", po.number], ["Order Date", DATE_STYLES["BUY"](po.date)],
               ["Deliver By", DATE_STYLES["BUY"](po.deliver_by)]])
    top = [two_col(party_block(BUYER), [Paragraph("PURCHASE ORDER", TITLE), Spacer(1, 4), head])]
    rows = [["Code", "Description", "Qty", "Unit", "Unit Price", "Amount"]]
    for l in po.lines:
        it = ITEMS[l.item]
        rows.append([it.sku, Paragraph(it.description, SMALL), qty(l.qty, it.unit), it.unit, money(l.price), money(l.amount)])
    sub = round(sum(l.amount for l in po.lines), 2)
    gst = round(sub * GST_RATE, 2)
    tot = kv([["Subtotal", money(sub)], ["GST", money(gst)], ["Order Total", money(sub + gst)]], widths=(35 * mm, 30 * mm))
    body = [Spacer(1, 8), Paragraph("<b>Supplier</b><br/>" + s.name + "<br/>" + "<br/>".join(s.address), NORMAL),
            Spacer(1, 10), grid(rows, [30 * mm, 62 * mm, 17 * mm, 11 * mm, 22 * mm, 24 * mm], align_right_from=4),
            Spacer(1, 8), Table([["", tot]], colWidths=[95 * mm, 80 * mm]),
            Spacer(1, 10), Paragraph("Please quote our PO number on all dockets and invoices.", SMALL)]
    return [top + body]


def statement_rows(supplier: str, as_at: date):
    """Statement lines from the supplier's own ledger (includes documents we never received)."""
    entries = []
    for inv in INVOICES + [UNSEEN]:
        if inv.supplier == supplier and inv.date <= as_at:
            entries.append((inv.date, inv.number, "Invoice", inv.total))
    for cn in CREDITS:
        if cn.supplier == supplier:
            entries.append((cn.date, cn.number, "Credit note", -cn.total))
    for p in PAYMENTS:
        if p.supplier == supplier:
            entries.append((p.date, p.number, "Payment received - thank you", -p.total))
    entries.sort()
    return entries


def statement_story(supplier: str, as_at: date):
    s = SUP[supplier]
    fmt = DATE_STYLES[supplier]
    entries = statement_rows(supplier, as_at)
    rows = [["Date", "Reference", "Details", "Debit", "Credit", "Balance"]]
    bal = 0.0
    rows.append(["", "", "Opening balance", "", "", money(0.0)])
    for d_, ref, det, amt in entries:
        bal = round(bal + amt, 2)
        rows.append([fmt(d_), ref, det, money(amt) if amt > 0 else "", money(-amt) if amt < 0 else "", money(bal)])
    head = kv([["Statement Date", fmt(as_at)], ["Account", "RIDG01"], ["Closing Balance", money(bal)]],
              widths=(32 * mm, 40 * mm))
    top = [two_col(party_block(s), [Paragraph("STATEMENT OF ACCOUNT", BIG), head])]
    body = [Spacer(1, 8), Paragraph("To: " + BUYER.name, NORMAL), Spacer(1, 8),
            grid(rows, [22 * mm, 28 * mm, 50 * mm, 22 * mm, 22 * mm, 24 * mm], align_right_from=3),
            Spacer(1, 10), Paragraph(f"Closing balance {money(bal)} is due. Please remit to BSB 063-000 Acct 1234 5678.", SMALL)]
    return [top + body], entries, bal


def remittance_story(p: Payment):
    s = SUP[p.supplier]
    invs = {i.number: i for i in INVOICES}
    rows = [["Invoice", "Invoice Date", "Invoice Amount", "Amount Paid"]]
    for ref, amt in p.allocations:
        i = invs[ref]
        rows.append([ref, DATE_STYLES["BUY"](i.date), money(i.total), money(amt)])
    head = kv([["Remittance No", p.number], ["Payment Date", DATE_STYLES["BUY"](p.date)], ["Method", "EFT"],
               ["Amount", money(p.total)]])
    top = [two_col(party_block(BUYER), [Paragraph("REMITTANCE ADVICE", BIG), head])]
    body = [Spacer(1, 8), Paragraph("Payee: " + s.name, NORMAL), Spacer(1, 8),
            grid(rows, [36 * mm, 30 * mm, 34 * mm, 34 * mm], align_right_from=2),
            Spacer(1, 8), Paragraph(f"Total paid: {money(p.total)}", NORMAL)]
    return [top + body]


def price_schedule_story():
    s = SUP["S1"]
    rows = [["Code", "Description", "Unit", "Price ex GST"]]
    for k, price in SCHEDULE_S1.items():
        it = ITEMS[k]
        rows.append([it.sku, it.description, it.unit, money(price)])
    head = kv([["Schedule", "PL-2026-H2"], ["Effective From", "01/07/2026"], ["Effective To", "31/12/2026"],
               ["Customer", "RIDG01"]])
    top = [two_col(party_block(s), [Paragraph("PRICE SCHEDULE", BIG), head])]
    body = [Spacer(1, 10), Paragraph("Contract pricing for " + BUYER.name + ". Prices exclude GST and delivery.", NORMAL),
            Spacer(1, 8), grid(rows, [36 * mm, 90 * mm, 16 * mm, 28 * mm], align_right_from=3)]
    return [top + body]


def sds_story():
    s = SUP["S3"]
    paras = [Paragraph("SAFETY DATA SHEET", BIG), Paragraph(f"Product: Inorganic zinc primer. Supplier: {s.name}.", NORMAL),
             Spacer(1, 6)]
    for h, t in [("1. Identification", "Recommended use: protective coating for structural steel. Emergency: 000."),
                 ("2. Hazard identification", "Flammable liquid category 3. Causes serious eye irritation."),
                 ("3. Composition", "Zinc dust 60-80%, ethyl silicate 10-20%, isopropanol 5-10%."),
                 ("4. First aid", "Eyes: rinse with water for 15 minutes. Skin: wash with soap and water."),
                 ("7. Handling and storage", "Store below 30C away from ignition sources. Keep container closed.")]:
        paras += [Paragraph(f"<b>{h}</b>", NORMAL), Paragraph(t, NORMAL), Spacer(1, 4)]
    return [paras]


# ============================================================== other formats
def write_docx_invoice(path: Path, inv: Invoice) -> None:
    import docx
    s = SUP[inv.supplier]
    fmt = DATE_STYLES[inv.supplier]
    d = docx.Document()
    d.add_heading(s.name, level=1)
    d.add_paragraph(", ".join(s.address) + f"\nABN: {s.abn}")
    d.add_heading("TAX INVOICE", level=2)
    d.add_paragraph(f"Invoice Number: {inv.number}\nInvoice Date: {fmt(inv.date)}\nPurchase Order: {inv.po}\n"
                    f"Terms: {inv.terms}\nCustomer: {BUYER.name}")
    t = d.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ["Product Code", "Description", "Qty", "Unit Price", "Total"]):
        c.text = h
    for l in inv.lines:
        it = ITEMS[l.item]
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text, r[4].text = it.sku, it.description, qty(l.qty, it.unit), money(l.price), money(l.amount)
    d.add_paragraph(f"Subtotal: {money(inv.subtotal)}\nGST: {money(inv.gst)}\nTotal Due: {money(inv.total)}")
    d.core_properties.author = "synthetic"
    d.save(str(path))


def write_ubl_invoice(path: Path, inv: Invoice) -> None:
    """A minimal Peppol PINT A-NZ style UBL 2.1 invoice."""
    s = SUP[inv.supplier]
    due = inv.date + timedelta(days=30)
    L = []
    for i, l in enumerate(inv.lines, 1):
        it = ITEMS[l.item]
        L.append(f"""  <cac:InvoiceLine>
    <cbc:ID>{i}</cbc:ID>
    <cbc:InvoicedQuantity unitCode="EA">{l.qty:g}</cbc:InvoicedQuantity>
    <cbc:LineExtensionAmount currencyID="AUD">{l.amount:.2f}</cbc:LineExtensionAmount>
    <cac:Item>
      <cbc:Name>{it.description}</cbc:Name>
      <cac:SellersItemIdentification><cbc:ID>{it.sku}</cbc:ID></cac:SellersItemIdentification>
      <cac:ClassifiedTaxCategory><cbc:ID>S</cbc:ID><cbc:Percent>10</cbc:Percent><cac:TaxScheme><cbc:ID>GST</cbc:ID></cac:TaxScheme></cac:ClassifiedTaxCategory>
    </cac:Item>
    <cac:Price><cbc:PriceAmount currencyID="AUD">{l.price:.2f}</cbc:PriceAmount></cac:Price>
  </cac:InvoiceLine>""")
    abn = s.abn.replace(" ", "")
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
  <cbc:CustomizationID>urn:peppol:pint:billing-1@aunz-1</cbc:CustomizationID>
  <cbc:ProfileID>urn:peppol:bis:billing</cbc:ProfileID>
  <cbc:ID>{inv.number}</cbc:ID>
  <cbc:IssueDate>{inv.date.isoformat()}</cbc:IssueDate>
  <cbc:DueDate>{due.isoformat()}</cbc:DueDate>
  <cbc:InvoiceTypeCode>380</cbc:InvoiceTypeCode>
  <cbc:DocumentCurrencyCode>AUD</cbc:DocumentCurrencyCode>
  <cac:OrderReference><cbc:ID>{inv.po}</cbc:ID></cac:OrderReference>
  <cac:AccountingSupplierParty><cac:Party>
    <cbc:EndpointID schemeID="0151">{abn}</cbc:EndpointID>
    <cac:PartyLegalEntity><cbc:RegistrationName>{s.name}</cbc:RegistrationName><cbc:CompanyID schemeID="0151">{abn}</cbc:CompanyID></cac:PartyLegalEntity>
  </cac:Party></cac:AccountingSupplierParty>
  <cac:AccountingCustomerParty><cac:Party>
    <cbc:EndpointID schemeID="0151">{BUYER.abn.replace(' ', '')}</cbc:EndpointID>
    <cac:PartyLegalEntity><cbc:RegistrationName>{BUYER.name}</cbc:RegistrationName></cac:PartyLegalEntity>
  </cac:Party></cac:AccountingCustomerParty>
  <cac:PaymentTerms><cbc:Note>{inv.terms}</cbc:Note></cac:PaymentTerms>
  <cac:TaxTotal>
    <cbc:TaxAmount currencyID="AUD">{inv.gst:.2f}</cbc:TaxAmount>
  </cac:TaxTotal>
  <cac:LegalMonetaryTotal>
    <cbc:LineExtensionAmount currencyID="AUD">{inv.subtotal:.2f}</cbc:LineExtensionAmount>
    <cbc:TaxExclusiveAmount currencyID="AUD">{inv.subtotal:.2f}</cbc:TaxExclusiveAmount>
    <cbc:TaxInclusiveAmount currencyID="AUD">{inv.total:.2f}</cbc:TaxInclusiveAmount>
    <cbc:PayableAmount currencyID="AUD">{inv.total:.2f}</cbc:PayableAmount>
  </cac:LegalMonetaryTotal>
{chr(10).join(L)}
</Invoice>
"""
    path.write_text(xml)


def write_xlsx_statement(path: Path, supplier: str, as_at: date):
    import openpyxl
    s = SUP[supplier]
    entries = statement_rows(supplier, as_at)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append([s.name])
    ws.append([f"ABN {s.abn}"])
    ws.append(["Statement of Account"])
    ws.append(["Customer", BUYER.name])
    ws.append(["Statement Date", as_at])
    ws.append([])
    ws.append(["Date", "Document", "Type", "Amount", "Running Balance"])
    bal = 0.0
    for d_, ref, det, amt in entries:
        bal = round(bal + amt, 2)
        ws.append([d_, ref, "INV" if amt > 0 and det == "Invoice" else ("CRN" if det == "Credit note" else "PMT"), amt, bal])
    ws.append([])
    ws.append(["", "", "Balance Due", bal])
    wb.save(str(path))
    return entries, bal


def write_xlsx_grn(path: Path, r: Receipt):
    import openpyxl
    po = next(p for p in POS if p.number == r.po)
    ordered = {l.item: l.qty for l in po.lines}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GRN"
    ws.append(["Goods Received Note", r.number])
    ws.append(["Supplier", SUP[r.supplier].name])
    ws.append(["PO Number", r.po])
    ws.append(["Received Date", r.date])
    ws.append([])
    ws.append(["Item Code", "Description", "Qty Ordered", "Qty Received", "UOM"])
    for l in r.lines:
        it = ITEMS[l.item]
        ws.append([it.sku, it.description, ordered[l.item], l.qty, it.unit])
    wb.save(str(path))


def write_csv_statement(path: Path, supplier: str, as_at: date):
    entries = statement_rows(supplier, as_at)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Supplier", SUP[supplier].name, "ABN", SUP[supplier].abn])
    w.writerow(["Customer", BUYER.name, "Statement date", as_at.isoformat()])
    w.writerow(["Trans Date", "Reference", "Description", "Debit", "Credit", "Balance"])
    bal = 0.0
    for d_, ref, det, amt in entries:
        bal = round(bal + amt, 2)
        w.writerow([d_.isoformat(), ref, det, f"{amt:.2f}" if amt > 0 else "", f"{-amt:.2f}" if amt < 0 else "", f"{bal:.2f}"])
    w.writerow(["", "", "Closing balance", "", "", f"{bal:.2f}"])
    path.write_text(buf.getvalue())
    return entries, bal


def write_eml(path: Path, sender: str, to: str, subject: str, body: str, when: date,
              attachments: list[tuple[str, bytes, str]] = ()) -> None:
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = sender, to, subject
    m["Date"] = when.strftime("%a, %d %b %Y 10:15:00 +1000")
    m["Message-ID"] = f"<{path.stem}@synthetic.example>"
    m.set_content(body)
    for name, data, mime in attachments:
        maintype, subtype = mime.split("/")
        m.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    path.write_bytes(bytes(m))


# ============================================================== image degradation
def render_pages(pdf: Path, dpi=150) -> list[Image.Image]:
    doc = pdfium.PdfDocument(str(pdf))
    return [doc[i].render(scale=dpi / 72).to_pil().convert("RGB") for i in range(len(doc))]


def photograph(img: Image.Image, rng: random.Random) -> Image.Image:
    """A phone photo of a paper document: tilt, uneven light, blur, JPEG."""
    img = img.rotate(rng.uniform(-3.5, 3.5), expand=True, fillcolor=(96, 88, 78))
    w, h = img.size
    shade = Image.linear_gradient("L").resize((w, h)).point(lambda v: 255 - v // 4)
    img = Image.composite(img, Image.new("RGB", (w, h), (170, 160, 140)), shade)
    img = img.filter(ImageFilter.GaussianBlur(0.9))
    img = img.resize((int(w * 0.8), int(h * 0.8)))
    return img


def scan(img: Image.Image, rng: random.Random) -> Image.Image:
    """An office scanner: greyscale, slight skew, speckle."""
    g = ImageOps.grayscale(img).rotate(rng.uniform(-1.2, 1.2), expand=False, fillcolor=255)
    px = g.load()
    w, h = g.size
    for _ in range(w * h // 900):
        px[rng.randrange(w), rng.randrange(h)] = rng.choice((0, 90, 160))
    return g.filter(ImageFilter.GaussianBlur(0.4))


# ============================================================== truth helpers
def inv_truth(inv: Invoice, doc_type="supplier_invoice"):
    s = SUP[inv.supplier]
    due = inv.date + timedelta(days=30 if "30" in inv.terms else 14)
    f = {"supplier_name": s.name, "supplier_abn": s.abn, "buyer_name": BUYER.name, "doc_number": inv.number,
         "date": inv.date.isoformat(), "po_reference": inv.po, "subtotal": inv.subtotal, "gst": inv.gst, "total": inv.total}
    if doc_type == "credit_note":
        f["original_invoice"] = inv.credit_for
    lines = [{"sku": ITEMS[l.item].sku, "description": ITEMS[l.item].description, "quantity": l.qty,
              "unit_price": l.price, "amount": l.amount} for l in inv.lines]
    return f, lines


def receipt_truth(r: Receipt):
    f = {"supplier_name": SUP[r.supplier].name, "doc_number": r.number, "date": r.date.isoformat(), "po_reference": r.po}
    if r.kind == "grn":
        f["buyer_name"] = BUYER.name
    return f, [{"sku": ITEMS[l.item].sku, "description": ITEMS[l.item].description, "quantity": l.qty} for l in r.lines]


def po_truth(po: PO):
    sub = round(sum(l.amount for l in po.lines), 2)
    gst = round(sub * GST_RATE, 2)
    f = {"supplier_name": SUP[po.supplier].name, "buyer_name": BUYER.name, "doc_number": po.number,
         "date": po.date.isoformat(), "delivery_date": po.deliver_by.isoformat(),
         "subtotal": sub, "gst": gst, "total": round(sub + gst, 2)}
    return f, [{"sku": ITEMS[l.item].sku, "description": ITEMS[l.item].description, "quantity": l.qty,
                "unit_price": l.price, "amount": l.amount} for l in po.lines]


def statement_truth(supplier, entries, bal, as_at):
    f = {"supplier_name": SUP[supplier].name, "supplier_abn": SUP[supplier].abn, "date": as_at.isoformat(),
         "closing_balance": bal}
    return f, [{"date": d_.isoformat(), "reference": ref, "amount": amt} for d_, ref, _, amt in entries]


def remit_truth(p: Payment):
    f = {"supplier_name": SUP[p.supplier].name, "buyer_name": BUYER.name, "doc_number": p.number,
         "date": p.date.isoformat(), "total": p.total}
    return f, [{"reference": ref, "amount": amt} for ref, amt in p.allocations]


# ============================================================== reconciliation truth
def reconciliation_truth():
    """The answers the pile should produce once every document is read correctly."""
    received = {}
    for r in RECEIPTS:
        for l in r.lines:
            received[(r.po, l.item)] = round(received.get((r.po, l.item), 0) + l.qty, 3)
    invoiced = {}
    for inv in INVOICES:
        if inv.matches_po:
            for l in inv.lines:
                invoiced[(inv.matches_po, l.item)] = round(invoiced.get((inv.matches_po, l.item), 0) + l.qty, 3)
    for cn in CREDITS:
        for l in cn.lines:
            invoiced[(cn.matches_po, l.item)] = round(invoiced.get((cn.matches_po, l.item), 0) - l.qty, 3)
    goods_owed, paperwork_owed, invoiced_not_received = [], [], []
    for po in POS:
        for l in po.lines:
            rec = received.get((po.number, l.item), 0)
            inv = invoiced.get((po.number, l.item), 0)
            sku = ITEMS[l.item].sku
            if l.qty - rec > 1e-9:
                entry = {"po": po.number, "sku": sku, "ordered": l.qty, "received": rec,
                         "outstanding": round(l.qty - rec, 3), "value": round((l.qty - rec) * l.price, 2)}
                if any(cn.matches_po == po.number and any(cl.item == l.item for cl in cn.lines) for cn in CREDITS):
                    # Open business-rule question for the client: does a credit for undelivered
                    # goods close the PO line, or is the shortfall still owed?
                    entry["credited_by"] = [cn.number for cn in CREDITS if cn.matches_po == po.number]
                goods_owed.append(entry)
            if rec - inv > 1e-9:
                paperwork_owed.append({"po": po.number, "sku": sku, "received": rec, "invoiced": inv,
                                       "outstanding": round(rec - inv, 3), "value": round((rec - inv) * l.price, 2)})
            if inv - rec > 1e-9:
                invoiced_not_received.append({"po": po.number, "sku": sku, "received": rec, "invoiced": inv,
                                              "excess": round(inv - rec, 3)})
    paid = {}
    for p in PAYMENTS:
        for ref, amt in p.allocations:
            paid[ref] = round(paid.get(ref, 0) + amt, 2)
    credited = {cn.credit_for: cn.total for cn in CREDITS}
    money_owed = []
    for inv in INVOICES:
        net = round(inv.total - credited.get(inv.number, 0), 2)
        out = round(net - paid.get(inv.number, 0), 2)
        if out > 0.004:
            money_owed.append({"invoice": inv.number, "supplier": SUP[inv.supplier].name, "invoice_total": inv.total,
                               "credits": credited.get(inv.number, 0), "paid": paid.get(inv.number, 0), "outstanding": out})
    exceptions = [
        {"bucket": "price_variance", "document": "BF-INV-30551", "po": "PO-4103", "sku": ITEMS["M16"].sku,
         "impact": 80.00, "detail": "billed 4.20 vs PO 4.00 on 400 ea"},
        {"bucket": "quantity_variance", "document": "HS-88302", "po": "PO-4102", "sku": ITEMS["RHS100"].sku,
         "impact": round(0.240 * 2310, 2), "detail": "billed 1.200 t, received 0.960 t", "resolved_by": "HS-CN-1190"},
        {"bucket": "no_matching_po", "document": "CC-10497", "impact": INVOICES[7].total},
        {"bucket": "missing_receipt", "document": "CC-10442", "po": "PO-4105", "impact": INVOICES[5].total},
        {"bucket": "duplicate_suspected", "document": "BF-INV-30551", "copies": 2},
        {"bucket": "duplicate_suspected", "document": "HS-88214", "copies": 2},
        {"bucket": "missing_document", "document": "HS-88377", "detail": "on Harbourline statement, never received",
         "impact": UNSEEN.total},
    ]
    links = {inv.number: inv.matches_po for inv in INVOICES}
    links.update({cn.number: cn.matches_po for cn in CREDITS})
    return {"goods_owed": goods_owed, "paperwork_owed": paperwork_owed,
            "invoiced_not_received": invoiced_not_received, "money_owed": money_owed,
            "exceptions": exceptions, "invoice_to_po": links,
            "note": "SG-77120 quotes no PO; it belongs to PO-4107 and should be found by fuzzy match."}


# ============================================================== main
def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    tmp = OUT / "_tmp"
    tmp.mkdir()
    docs = []
    rng = random.Random(7)

    def add(files, doc_type, fmt, fields, lines, pages=None, doc_id=None):
        docs.append({"id": doc_id or fields.get("doc_number") or files[0], "files": files, "pages": pages,
                     "doc_type": doc_type, "format": fmt, "fields": fields, "lines": lines})

    inv = {i.number: i for i in INVOICES}
    rec = {r.number: r for r in RECEIPTS}

    # --- purchase orders: born-digital PDFs, one exported to a spreadsheet
    for po in POS:
        if po.number == "PO-4108":
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Purchase Order", po.number])
            ws.append(["Supplier", SUP[po.supplier].name])
            ws.append(["Order Date", po.date])
            ws.append(["Deliver By", po.deliver_by])
            ws.append([])
            ws.append(["Code", "Description", "Qty", "Unit", "Unit Price", "Amount"])
            for l in po.lines:
                it = ITEMS[l.item]
                ws.append([it.sku, it.description, l.qty, it.unit, l.price, l.amount])
            sub = sum(l.amount for l in po.lines)
            ws.append(["", "", "", "", "Subtotal", sub])
            ws.append(["", "", "", "", "GST", round(sub * 0.1, 2)])
            ws.append(["", "", "", "", "Total", round(sub * 1.1, 2)])
            ws.append(["Issued by", BUYER.name])
            name = f"{po.number}.xlsx"
            wb.save(str(OUT / name))
            f, l_ = po_truth(po)
            add([name], "purchase_order", "xlsx", f, l_)
        else:
            name = f"{po.number}.pdf"
            build_pdf(OUT / name, po_story(po))
            f, l_ = po_truth(po)
            add([name], "purchase_order", "pdf_text", f, l_)

    # --- invoices
    build_pdf(OUT / "HS-88214.pdf", invoice_story(inv["HS-88214"]))
    add(["HS-88214.pdf"], "supplier_invoice", "pdf_text", *inv_truth(inv["HS-88214"]))
    # phone photo of the same invoice: a duplicate arriving by another route
    photo = photograph(render_pages(OUT / "HS-88214.pdf")[0], rng)
    photo.save(OUT / "IMG_20260808_091412.jpg", quality=62)
    add(["IMG_20260808_091412.jpg"], "supplier_invoice", "image_photo", *inv_truth(inv["HS-88214"]),
        doc_id="HS-88214 (photo)")

    build_pdf(OUT / "HS-88302.pdf", invoice_story(inv["HS-88302"], page_split=1))
    add(["HS-88302.pdf"], "supplier_invoice", "pdf_text_multipage", *inv_truth(inv["HS-88302"]), pages={"HS-88302.pdf": [1, 2]})

    build_pdf(OUT / "HS-CN-1190.pdf", invoice_story(CREDITS[0], credit=True))
    add(["HS-CN-1190.pdf"], "credit_note", "pdf_text", *inv_truth(CREDITS[0], "credit_note"))

    # Boltmaster's pack: docket + invoice + docket in one born-digital PDF
    pack = OUT / "Boltmaster_documents_Aug.pdf"
    build_pdf(pack, receipt_story(rec["BF-20931"]) + invoice_story(inv["BF-INV-30551"]) + receipt_story(rec["BF-21007"]))
    add([pack.name], "goods_receipt", "pdf_text_pack", *receipt_truth(rec["BF-20931"]), pages={pack.name: [1]})
    add([pack.name], "supplier_invoice", "pdf_text_pack", *inv_truth(inv["BF-INV-30551"]), pages={pack.name: [2]})
    add([pack.name], "goods_receipt", "pdf_text_pack", *receipt_truth(rec["BF-21007"]), pages={pack.name: [3]})

    # the same invoice re-sent as a "COPY" by email: different bytes, same document
    build_pdf(tmp / "BF-INV-30551.pdf", invoice_story(inv["BF-INV-30551"], copy_stamp=True))
    write_eml(OUT / "fwd_copy_invoice_BF-INV-30551.eml", SUP["S2"].email, BUYER.email,
              "Copy invoice BF-INV-30551", "Hi,\n\nAs requested please find attached a copy of invoice BF-INV-30551.\n\nRegards\nBoltmaster Accounts",
              D(27), [("BF-INV-30551.pdf", (tmp / "BF-INV-30551.pdf").read_bytes(), "application/pdf")])
    add(["fwd_copy_invoice_BF-INV-30551.eml#BF-INV-30551.pdf"], "supplier_invoice", "eml_attachment",
        *inv_truth(inv["BF-INV-30551"]), doc_id="BF-INV-30551 (emailed copy)")

    write_ubl_invoice(OUT / "BF-INV-30602.xml", inv["BF-INV-30602"])
    add(["BF-INV-30602.xml"], "supplier_invoice", "ubl_xml", *inv_truth(inv["BF-INV-30602"]))

    # two-page invoice where page 2 arrived as its own file
    pages = invoice_story(inv["BF-INV-30688"], page_split=1)
    # only one line: split puts line on p1 and totals on p2
    build_pdf(OUT / "BF-INV-30688.pdf", [pages[0]])
    build_pdf(OUT / "BF-INV-30688_page2.pdf", [pages[1]])
    add(["BF-INV-30688.pdf", "BF-INV-30688_page2.pdf"], "supplier_invoice", "pdf_text_split_files",
        *inv_truth(inv["BF-INV-30688"]))

    write_docx_invoice(OUT / "CC-10442.docx", inv["CC-10442"])
    add(["CC-10442.docx"], "supplier_invoice", "docx", *inv_truth(inv["CC-10442"]))

    build_pdf(OUT / "SG-77120.pdf", invoice_story(inv["SG-77120"]))
    add(["SG-77120.pdf"], "supplier_invoice", "pdf_text", *inv_truth(inv["SG-77120"]))

    build_pdf(OUT / "CC-10497.pdf", invoice_story(inv["CC-10497"]))
    add(["CC-10497.pdf"], "supplier_invoice", "pdf_text", *inv_truth(inv["CC-10497"]))

    # --- receipts
    build_pdf(tmp / "DD-55120.pdf", receipt_story(rec["DD-55120"]))
    photograph(render_pages(tmp / "DD-55120.pdf")[0], rng).save(OUT / "IMG_20260807_143055.jpg", quality=58)
    add(["IMG_20260807_143055.jpg"], "goods_receipt", "image_photo", *receipt_truth(rec["DD-55120"]))

    build_pdf(OUT / "GRN-7002.pdf", receipt_story(rec["GRN-7002"]))
    add(["GRN-7002.pdf"], "goods_receipt", "pdf_text", *receipt_truth(rec["GRN-7002"]))

    write_xlsx_grn(OUT / "GRN-7006.xlsx", rec["GRN-7006"])
    add(["GRN-7006.xlsx"], "goods_receipt", "xlsx", *receipt_truth(rec["GRN-7006"]))

    # front-desk scanner batch: two dockets in one image-only PDF
    build_pdf(tmp / "scanpack.pdf", receipt_story(rec["BF-21088"]) + receipt_story(rec["SG-3301"]))
    imgs = [scan(p, rng) for p in render_pages(tmp / "scanpack.pdf", dpi=150)]
    imgs[0].save(OUT / "scan_frontdesk_0826.pdf", save_all=True, append_images=imgs[1:], resolution=150)
    add(["scan_frontdesk_0826.pdf"], "goods_receipt", "pdf_scan", *receipt_truth(rec["BF-21088"]), pages={"scan_frontdesk_0826.pdf": [1]})
    add(["scan_frontdesk_0826.pdf"], "goods_receipt", "pdf_scan", *receipt_truth(rec["SG-3301"]), pages={"scan_frontdesk_0826.pdf": [2]})

    # --- statements
    story, entries, bal = statement_story("S1", D(31))
    build_pdf(OUT / "Harbourline_statement_Aug26.pdf", story)
    add(["Harbourline_statement_Aug26.pdf"], "statement", "pdf_text", *statement_truth("S1", entries, bal, D(31)),
        doc_id="statement S1")
    entries, bal = write_xlsx_statement(OUT / "Boltmaster_statement_2026-08.xlsx", "S2", D(31))
    add(["Boltmaster_statement_2026-08.xlsx"], "statement", "xlsx", *statement_truth("S2", entries, bal, D(31)),
        doc_id="statement S2")
    entries, bal = write_csv_statement(OUT / "SGWS_statement_20260831.csv", "S4", D(31))
    add(["SGWS_statement_20260831.csv"], "statement", "csv", *statement_truth("S4", entries, bal, D(31)),
        doc_id="statement S4")

    # --- remittances
    build_pdf(tmp / "RA-9001.pdf", remittance_story(PAYMENTS[0]))
    write_eml(OUT / "remittance_RA-9001.eml", BUYER.email, SUP["S1"].email, "Remittance advice RA-9001",
              "Please find our remittance advice attached.\n\nRidgeline Accounts", D(29),
              [("RA-9001.pdf", (tmp / "RA-9001.pdf").read_bytes(), "application/pdf")])
    add(["remittance_RA-9001.eml#RA-9001.pdf"], "remittance", "eml_attachment", *remit_truth(PAYMENTS[0]))

    build_pdf(OUT / "RA-9002.pdf", remittance_story(PAYMENTS[1]))
    add(["RA-9002.pdf"], "remittance", "pdf_text", *remit_truth(PAYMENTS[1]))

    p = PAYMENTS[2]
    body = (f"REMITTANCE ADVICE\nRemittance No: {p.number}\nPayment Date: {DATE_STYLES['BUY'](p.date)}\n"
            f"Payer: {BUYER.name}\nPayee: {SUP[p.supplier].name}\n\nInvoice    Amount Paid\n" +
            "\n".join(f"{ref}    {money(a)}" for ref, a in p.allocations) +
            f"\n\nTotal paid: {money(p.total)}\nPart payment, balance to follow.\n\nRidgeline Accounts")
    write_eml(OUT / "remittance_RA-9003.eml", BUYER.email, SUP["S4"].email, f"Payment {p.number}", body, D(29))
    add(["remittance_RA-9003.eml"], "remittance", "eml_body", *remit_truth(p))

    # --- price schedule and a document that is none of the eight types
    build_pdf(OUT / "Harbourline_price_schedule_H2-2026.pdf", price_schedule_story())
    add(["Harbourline_price_schedule_H2-2026.pdf"], "price_schedule", "pdf_text",
        {"supplier_name": SUP["S1"].name, "doc_number": "PL-2026-H2", "effective_from": "2026-07-01", "effective_to": "2026-12-31"},
        [{"sku": ITEMS[k].sku, "description": ITEMS[k].description, "unit_price": v} for k, v in SCHEDULE_S1.items()])
    build_pdf(OUT / "SDS_zinc_primer.pdf", sds_story())
    add(["SDS_zinc_primer.pdf"], "unknown", "pdf_text", {}, [], doc_id="SDS")

    shutil.rmtree(tmp)
    truth = {"scenario": "Ridgeline Steel Fabrications, August 2026 (synthetic; all parties invented)",
             "buyer": {"name": BUYER.name, "abn": BUYER.abn},
             "suppliers": {k: {"name": v.name, "abn": v.abn} for k, v in SUP.items()},
             "documents": docs, "reconciliation": reconciliation_truth()}
    (OUT / "truth.json").write_text(json.dumps(truth, indent=2, default=str))
    files = sorted(p.name for p in OUT.iterdir() if p.name != "truth.json")
    print(f"{len(files)} files, {len(docs)} documents -> {OUT}")


if __name__ == "__main__":
    main()
