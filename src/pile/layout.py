"""One shape for every readable format: a grid of rows of cells.

Spreadsheets are grids already. Word files become paragraphs and table rows.
Plain text (email bodies) splits on runs of two or more spaces. PDF text layers
(pdf_text.py) group words into rows by position and cells by horizontal gaps.
Extraction then works on the grid without caring where it came from.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any


@dataclass
class Row:
    cells: list[Any]
    page: int = 1
    y: float | None = None
    xs: list[float] | None = None          # left x of each cell, PDFs only

    @property
    def texts(self) -> list[str]:
        return [cell_text(c) for c in self.cells]

    @property
    def is_blank(self) -> bool:
        return not any(t.strip() for t in self.texts)

    def __str__(self) -> str:
        return " | ".join(t for t in self.texts if t)


@dataclass
class Grid:
    rows: list[Row] = field(default_factory=list)
    pages: int = 1
    source: str = ""

    @property
    def text(self) -> str:
        return "\n".join("  ".join(t for t in r.texts if t) for r in self.rows)

    def page_rows(self, page: int) -> list[Row]:
        return [r for r in self.rows if r.page == page]


def cell_text(v: Any) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return f"{v:.2f}" if abs(v - round(v, 2)) < 1e-9 and not float(v).is_integer() else f"{v:g}"
    return str(v).strip()


def _trim(values: list[Any]) -> list[Any]:
    vals = list(values)
    while vals and (vals[-1] is None or cell_text(vals[-1]) == ""):
        vals.pop()
    return vals


def from_xlsx(data: bytes, source: str = "") -> Grid:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows: list[Row] = []
    for i, ws in enumerate(wb.worksheets, 1):
        for vals in ws.iter_rows(values_only=True):
            rows.append(Row(_trim(list(vals)), page=i))
    return Grid(rows, pages=len(wb.worksheets), source=source)


def from_csv(data: bytes, source: str = "") -> Grid:
    text = data.decode("utf-8-sig", errors="replace")
    head = text[:4096]
    delim = max(",;\t|", key=head.count) if head.strip() else ","
    return Grid([Row(_trim(r)) for r in csv.reader(io.StringIO(text), delimiter=delim)], source=source)


def from_docx(data: bytes, source: str = "") -> Grid:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(io.BytesIO(data))
    rows: list[Row] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            for line in Paragraph(child, d).text.splitlines():
                rows.append(Row(split_line(line)))
        elif tag == "tbl":
            for tr in Table(child, d).rows:
                rows.append(Row(_trim([c.text.strip() for c in tr.cells])))
            rows.append(Row([]))
    return Grid(rows, source=source)


def split_line(line: str) -> list[str]:
    """Cells in a line of plain text: separated by tabs or two-plus spaces."""
    parts = [p.strip() for p in re.split(r"\t+|\s{2,}", line.strip())]
    return [p for p in parts if p]


def from_text(data: bytes | str, source: str = "") -> Grid:
    text = data.decode("utf-8", errors="replace") if isinstance(data, bytes) else data
    return Grid([Row(split_line(l)) for l in text.splitlines()], source=source)
